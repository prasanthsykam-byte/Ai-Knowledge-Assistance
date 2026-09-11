from __future__ import annotations

import csv
import json
from io import BytesIO, StringIO
from pathlib import Path

from docx import Document
from langchain_core.documents import Document as LangChainDocument
from openpyxl import load_workbook
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".pdf", ".docx", ".xlsx"}
MAX_FILE_SIZE = 25 * 1024 * 1024


def extract_text(uploaded_file) -> str:
    suffix = Path(uploaded_file.name).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {suffix or 'unknown'}")
    content = uploaded_file.getvalue()
    if len(content) > MAX_FILE_SIZE:
        raise ValueError("Each file must be smaller than 25 MB.")
    if suffix in {".txt", ".md"}:
        return content.decode("utf-8", errors="replace")
    if suffix == ".csv":
        rows = csv.reader(StringIO(content.decode("utf-8", errors="replace")))
        return "\n".join(", ".join(row) for row in rows)
    if suffix == ".json":
        return json.dumps(json.loads(content.decode("utf-8", errors="replace")), indent=2)
    if suffix == ".pdf":
        return "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(content)).pages)
    if suffix == ".docx":
        document = Document(BytesIO(content))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheets = []
    for worksheet in workbook.worksheets:
        rows = [", ".join("" if value is None else str(value) for value in row) for row in worksheet.iter_rows(values_only=True)]
        sheets.append(f"Sheet: {worksheet.title}\n" + "\n".join(rows))
    return "\n\n".join(sheets)


def to_document(text: str, name: str, document_id: int, user_id: int) -> LangChainDocument:
    return LangChainDocument(
        page_content=text,
        metadata={"source": name, "document_id": str(document_id), "user_id": str(user_id)},
    )
